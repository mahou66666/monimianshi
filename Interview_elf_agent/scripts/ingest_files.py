#!/usr/bin/env python3

# Interview_elf_agent/scripts/ingest_files.py

import argparse
import hashlib
import uuid
import json
import os
import sys
from typing import Dict, Iterable, List, Optional, Tuple

import pandas as pd

# 确保能导入 app 模块
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from qdrant_client import QdrantClient
from qdrant_client.http import models

from app.core.config import settings
from app.utils.model_loader import get_embedding_model


KB_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data_source")


def _stable_hash(text: str) -> str:
    return hashlib.md5(text.encode("utf-8")).hexdigest()


def _make_doc_id(source: str, content: str, explicit_id: Optional[str]) -> str:
    if explicit_id:
        return str(explicit_id)
    return _stable_hash(f"{source}\n{content}")


def _make_point_id(doc_id: str, chunk_id: int) -> str:
    return str(uuid.uuid5(uuid.NAMESPACE_URL, f"{doc_id}:{chunk_id}"))


def _record(
    text: str,
    source: str,
    data_type: str,
    explicit_id: Optional[str] = None,
    extra: Optional[Dict] = None,
    is_qa: bool = False,
) -> Dict:
    metadata = {"source": source, "type": data_type}
    if extra:
        metadata.update(extra)
    return {
        "text": text,
        "metadata": metadata,
        "doc_id": _make_doc_id(source, text, explicit_id),
        "is_qa": is_qa,
    }


def load_docx(file_path: str) -> List[Dict]:
    if DocxDocument is None:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")
    docs = []
    doc = DocxDocument(file_path)
    full_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
    if not full_text:
        return docs
    docs.append(
        _record(
            text=full_text,
            source=os.path.basename(file_path),
            data_type="docx",
        )
    )
    return docs


def load_excel(file_path: str) -> List[Dict]:
    docs = []
    df = pd.read_excel(file_path).fillna("")
    for _, row in df.iterrows():
        content_parts = [f"{col}: {val}" for col, val in row.items() if str(val).strip()]
        content = "\n".join(content_parts)
        if content:
            docs.append(
                _record(
                    text=content,
                    source=os.path.basename(file_path),
                    data_type="excel",
                )
            )
    return docs


def load_csv(file_path: str) -> List[Dict]:
    docs = []
    df = pd.read_csv(file_path).fillna("")
    cols = {c.lower(): c for c in df.columns}
    has_qa = "question" in cols and "answer" in cols
    for _, row in df.iterrows():
        if has_qa:
            question = str(row[cols["question"]]).strip()
            answer = str(row[cols["answer"]]).strip()
            if not question or not answer:
                continue
            text = f"Q: {question}\nA: {answer}"
            extra = {
                "question": question,
                "answer": answer,
                "major": str(row.get(cols.get("major"), "")).strip(),
                "topic": str(row.get(cols.get("topic"), "")).strip(),
                "subtopic": str(row.get(cols.get("subtopic"), "")).strip(),
                "keywords": row.get(cols.get("keywords"), []),
                "updated_at": str(row.get(cols.get("updated_at"), "")).strip(),
            }
            docs.append(
                _record(
                    text=text,
                    source=os.path.basename(file_path),
                    data_type="qa",
                    explicit_id=str(row.get(cols.get("id"), "")).strip() or None,
                    extra=extra,
                    is_qa=True,
                )
            )
        else:
            content_parts = [f"{col}: {val}" for col, val in row.items() if str(val).strip()]
            content = "\n".join(content_parts)
            if content:
                docs.append(
                    _record(
                        text=content,
                        source=os.path.basename(file_path),
                        data_type="csv",
                    )
                )
    return docs


def load_jsonl(file_path: str) -> List[Dict]:
    docs = []
    with open(file_path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            question = str(rec.get("question", "")).strip()
            answer = str(rec.get("answer", "")).strip()
            if question and answer:
                text = f"Q: {question}\nA: {answer}"
                extra = {
                    "question": question,
                    "answer": answer,
                    "major": rec.get("major", ""),
                    "topic": rec.get("topic", ""),
                    "subtopic": rec.get("subtopic", ""),
                    "keywords": rec.get("keywords", []),
                    "updated_at": rec.get("updated_at", ""),
                }
                docs.append(
                    _record(
                        text=text,
                        source=rec.get("source") or os.path.basename(file_path),
                        data_type="qa",
                        explicit_id=rec.get("id"),
                        extra=extra,
                        is_qa=True,
                    )
                )
            else:
                content = str(rec.get("text") or rec.get("content") or "").strip()
                if content:
                    docs.append(
                        _record(
                            text=content,
                            source=rec.get("source") or os.path.basename(file_path),
                            data_type="jsonl",
                        )
                    )
    return docs


def iter_files(input_path: str) -> Iterable[str]:
    if os.path.isfile(input_path):
        yield input_path
        return
    for root, _, files in os.walk(input_path):
        for name in files:
            if name.startswith("~$") or name.startswith("."):
                continue
            yield os.path.join(root, name)


def _ensure_collection(
    client: QdrantClient, collection_name: str, vector_size: int, recreate: bool
) -> None:
    if client.collection_exists(collection_name):
        info = client.get_collection(collection_name)
        vectors = info.config.params.vectors
        if isinstance(vectors, dict):
            existing_dim = next(iter(vectors.values())).size
        else:
            existing_dim = vectors.size
        if existing_dim != vector_size:
            if not recreate:
                raise RuntimeError(
                    f"Collection dim mismatch: existing={existing_dim}, new={vector_size}. "
                    "Use --recreate to rebuild."
                )
            print(f"♻️  Recreating collection: {collection_name}")
            client.delete_collection(collection_name)
        else:
            return

    client.create_collection(
        collection_name=collection_name,
        vectors_config=models.VectorParams(size=vector_size, distance=models.Distance.COSINE),
    )

    for field_name in ("major", "topic", "subtopic", "source", "type"):
        try:
            client.create_payload_index(
                collection_name=collection_name,
                field_name=field_name,
                field_schema=models.PayloadSchemaType.KEYWORD,
            )
        except Exception:
            pass


def ingest_data(
    input_path: str,
    source: str,
    mode: str,
    batch_size: int,
    chunk_size: int,
    chunk_overlap: int,
    recreate: bool,
) -> None:
    embeddings = get_embedding_model()
    print(f"🔌 连接 Qdrant ({settings.QDRANT_HOST}:{settings.QDRANT_PORT})...")
    client = QdrantClient(host=settings.QDRANT_HOST, port=settings.QDRANT_PORT)

    if not os.path.exists(input_path):
        raise FileNotFoundError(f"Input path not found: {input_path}")

    print("🔍 探测 embedding 维度...")
    sample_vec = embeddings.embed_query("dimension check")
    vector_size = len(sample_vec)
    _ensure_collection(client, settings.COLLECTION_NAME, vector_size, recreate)

    raw_records: List[Dict] = []
    for file_path in iter_files(input_path):
        lower = file_path.lower()
        if source != "auto":
            if source == "docx" and not lower.endswith(".docx"):
                continue
            if source == "excel" and not (lower.endswith(".xlsx") or lower.endswith(".xls")):
                continue
            if source == "csv" and not lower.endswith(".csv"):
                continue
            if source == "jsonl" and not lower.endswith(".jsonl"):
                continue

        if lower.endswith(".docx"):
            raw_records.extend(load_docx(file_path))
        elif lower.endswith(".xlsx") or lower.endswith(".xls"):
            raw_records.extend(load_excel(file_path))
        elif lower.endswith(".csv"):
            raw_records.extend(load_csv(file_path))
        elif lower.endswith(".jsonl"):
            raw_records.extend(load_jsonl(file_path))

    if not raw_records:
        print("❌ 未提取到有效数据。")
        return

    print(f"📄 原始记录数量: {len(raw_records)}")

    # 去重（同一批次）
    seen = set()
    unique_records = []
    for rec in raw_records:
        h = _stable_hash(rec["text"])
        if h in seen:
            continue
        seen.add(h)
        unique_records.append(rec)

    print(f"✅ 去重后记录数: {len(unique_records)}")

    # 切分
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap,
        separators=["\n\n", "\n", "。", "！", "？", " ", ""],
    )

    chunks: List[Dict] = []
    counters: Dict[str, int] = {}
    if unique_records:
        qa_records = [r for r in unique_records if r["is_qa"]]
        text_records = [r for r in unique_records if not r["is_qa"]]

        for rec in qa_records:
            doc_id = rec["doc_id"]
            chunk_id = 0
            payload = dict(rec["metadata"])
            payload.update({"doc_id": doc_id, "chunk_id": chunk_id, "page_content": rec["text"]})
            chunks.append(
                {
                    "id": _make_point_id(doc_id, chunk_id),
                    "text": rec["text"],
                    "payload": payload,
                }
            )

        if text_records:
            docs = [
                Document(page_content=r["text"], metadata={**r["metadata"], "doc_id": r["doc_id"]})
                for r in text_records
            ]
            split_docs = text_splitter.split_documents(docs)
            for doc in split_docs:
                doc_id = doc.metadata.get("doc_id") or _stable_hash(doc.page_content)
                counters.setdefault(doc_id, 0)
                chunk_id = counters[doc_id]
                counters[doc_id] += 1
                payload = dict(doc.metadata)
                payload.update({"doc_id": doc_id, "chunk_id": chunk_id, "page_content": doc.page_content})
                chunks.append(
                    {
                        "id": _make_point_id(doc_id, chunk_id),
                        "text": doc.page_content,
                        "payload": payload,
                    }
                )

    print(f"✅ 切分完成！生成 {len(chunks)} 个知识片段。")

    # 入库（增量 upsert）
    print(f"💾 开始向量化并写入数据库 (mode={mode})...")
    for i in range(0, len(chunks), batch_size):
        batch = chunks[i : i + batch_size]
        texts = [b["text"] for b in batch]
        vectors = embeddings.embed_documents(texts)
        points = [
            models.PointStruct(id=batch[idx]["id"], vector=vectors[idx], payload=batch[idx]["payload"])
            for idx in range(len(batch))
        ]
        client.upsert(collection_name=settings.COLLECTION_NAME, points=points)
        print(f"  -> 进度: {min(i + batch_size, len(chunks))}/{len(chunks)}")

    print("🎉 入库完成！")


def main() -> None:
    parser = argparse.ArgumentParser(description="Ingest files into Qdrant (incremental).")
    parser.add_argument("--input", default=KB_DIR, help="Input file or directory")
    parser.add_argument(
        "--source",
        choices=["auto", "docx", "excel", "jsonl", "csv"],
        default="auto",
        help="Data source type",
    )
    parser.add_argument(
        "--mode",
        choices=["full", "incremental"],
        default="incremental",
        help="Ingest mode (both use upsert; no collection drop by default)",
    )
    parser.add_argument("--batch-size", type=int, default=64)
    parser.add_argument("--chunk-size", type=int, default=700)
    parser.add_argument("--chunk-overlap", type=int, default=100)
    parser.add_argument(
        "--recreate", action="store_true", help="Recreate collection if dim mismatch."
    )
    args = parser.parse_args()
    ingest_data(
        input_path=args.input,
        source=args.source,
        mode=args.mode,
        batch_size=args.batch_size,
        chunk_size=args.chunk_size,
        chunk_overlap=args.chunk_overlap,
        recreate=args.recreate,
    )


if __name__ == "__main__":
    main()
